from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from .calculations import (
    acceleration_to_register,
    calculate_composite_motion,
    calculate_motion,
    derive_axis,
    distance_to_register,
    tvmax_to_register,
    velocity_to_register,
)
from .models import AppState, Axis, CompositeAction, CompositeStepItem, DistanceCase, MotionAction


MECHANISM_NAMES = {
    "linear_belt": "步进直线",
    "rotary": "步进旋转",
    "leadscrew": "步进丝杆",
    "custom": "自定义换算",
    "timer": "定时器",
}


def _fmt_time(value: float | int | str) -> float | str:
    if value == "":
        return ""
    return round(float(value), 3)


def _action_axis(state: AppState, action: MotionAction) -> Axis | None:
    return next((axis for axis in state.axes if axis.id == action.axis_id), None)


def _axis_index(state: AppState, axis: Axis) -> int | str:
    return next((idx for idx, item in enumerate(state.axes, start=1) if item.id == axis.id), "")


def _distance_lookup(state: AppState) -> dict[str, tuple[Axis, MotionAction, DistanceCase]]:
    lookup: dict[str, tuple[Axis, MotionAction, DistanceCase]] = {}
    for action in state.actions:
        axis = _action_axis(state, action)
        if not axis:
            continue
        for distance in action.distances:
            lookup[distance.id] = (axis, action, distance)
    return lookup


def _move_label(lookup: dict[str, tuple[Axis, MotionAction, DistanceCase]], distance_id: str) -> str:
    if not distance_id:
        return ""
    item = lookup.get(distance_id)
    if not item:
        return ""
    axis, _action, distance = item
    note = distance.note.strip() or f"{distance.distance:g}"
    if axis.normalized_type() == "timer":
        return f"{axis.name} / {note} / {distance.duration_s:g}s"
    return f"{axis.name} / {note} / {distance.distance:g}"


def _point_name(axis: Axis, point_id: str) -> str:
    point = next((item for item in axis.points if item.id == point_id), None)
    return point.name if point else ""


def _motion_distance(distance: DistanceCase) -> DistanceCase:
    return DistanceCase(
        id=distance.id,
        distance=abs(distance.distance),
        note=distance.note,
        start_point_id=distance.start_point_id,
        end_point_id=distance.end_point_id,
    )


def _auto_width(ws) -> None:
    for column in ws.columns:
        max_length = 0
        letter = column[0].column_letter
        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_length + 2, 10), 42)


def _velocity_reg(axis: Axis, state: AppState, value: float) -> int:
    return velocity_to_register(axis, state.fclk_hz, value).value


def _accel_reg(axis: Axis, state: AppState, value: float) -> int:
    return acceleration_to_register(axis, state.fclk_hz, value).value


def _composite_number(action: CompositeAction, index: int) -> str:
    return action.number.strip() or str(index)


def export_excel(path: str | Path, state: AppState) -> None:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "项目与单轴参数"
    ws1.append(["项目信息", ""])
    ws1.append(["项目名称", state.project_name])
    ws1.append(["fCLK Hz", state.fclk_hz])
    ws1.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws1.append([])
    ws1.append(["运动轴列表", ""])
    ws1.append([
        "轴编号", "名称", "结构类型", "步距角", "微步", "减速比",
        "同步轮齿数", "同步轮齿距 mm", "丝杆导程 mm/rev",
        "电机每圈微步数", "每圈位移", "单位", "microsteps/unit",
    ])
    for idx, axis in enumerate(state.axes, start=1):
        try:
            derived = derive_axis(axis)
            derived_values = [derived.count, derived.unit_per_motor_rev, derived.unit_name, derived.microsteps_per_unit]
        except Exception:
            derived_values = ["", "", "", ""]
        ws1.append([
            idx,
            axis.name,
            MECHANISM_NAMES.get(axis.normalized_type(), axis.normalized_type()),
            axis.motor_step_angle,
            axis.microstep,
            axis.gear_ratio,
            axis.pulley_teeth,
            axis.pulley_pitch_mm,
            axis.lead_mm_per_rev,
            *derived_values,
        ])

    ws1.append([])
    ws1.append(["单轴参数（寄存器值）", ""])
    ws1.append(["轴编号", "名称", "VMAX", "TMAX", "V1", "V2", "A1", "A2", "AMAX", "DMAX", "D2", "D1"])
    for action in state.actions:
        axis = _action_axis(state, action)
        if not axis:
            continue
        if axis.normalized_type() == "timer":
            ws1.append([_axis_index(state, axis), axis.name, "", "", "", "", "", "", "", "", "", ""])
            continue
        params = action.params
        ws1.append([
            _axis_index(state, axis),
            axis.name,
            _velocity_reg(axis, state, params.vmax),
            tvmax_to_register(state.fclk_hz, params.tvmax_ms).value,
            _velocity_reg(axis, state, params.v1),
            _velocity_reg(axis, state, params.v2),
            _accel_reg(axis, state, params.a1),
            _accel_reg(axis, state, params.a2),
            _accel_reg(axis, state, params.amax),
            _accel_reg(axis, state, params.dmax),
            _accel_reg(axis, state, params.d2),
            _accel_reg(axis, state, params.d1),
        ])

    ws2 = wb.create_sheet("距离坐标")
    ws2.append(["轴名称", "距离备注", "起始点", "结束点", "距离 mm", "X_TARGET", "曲线", "峰值速度寄存器", "运动时间 s"])
    for action in state.actions:
        axis = _action_axis(state, action)
        if not axis:
            continue
        for distance in action.distances:
            if axis.normalized_type() == "timer":
                ws2.append([axis.name, distance.note, "", "", "", "", "定时器", "", _fmt_time(distance.duration_s)])
                continue
            try:
                result = calculate_motion(axis, state.fclk_hz, action.params, _motion_distance(distance))
                profile = result.profile_type
                vpeak = _velocity_reg(axis, state, result.vpeak)
                total = _fmt_time(result.total_time)
                xtarget = distance_to_register(axis, distance.distance).value
            except Exception:
                profile = ""
                vpeak = ""
                total = ""
                xtarget = distance_to_register(axis, distance.distance).value
            ws2.append([axis.name, distance.note, _point_name(axis, distance.start_point_id), _point_name(axis, distance.end_point_id), distance.distance, xtarget, profile, vpeak, total])

    ws3 = wb.create_sheet("一级动作时间")
    ws3.append(["编号", "动作名称", "时间 s"])
    composite_results = {}
    for idx, action in enumerate(state.composite_actions, start=1):
        try:
            result = calculate_composite_motion(state.axes, state.actions, action.steps, state.fclk_hz)
            total = _fmt_time(result.total_time)
            composite_results[action.id] = result
        except Exception:
            total = ""
            composite_results[action.id] = None
        ws3.append([_composite_number(action, idx), action.name, total])

    ws4 = wb.create_sheet("组合STEP列表")
    ws4.append([
        "组合编号", "组合名称", "STEP", "备注", "动作轴", "距离或延时",
        "动作时间 s", "STEP时间 s", "组合动作时间 s",
    ])
    lookup = _distance_lookup(state)
    row = 2
    for action_idx, action in enumerate(state.composite_actions, start=1):
        result = composite_results.get(action.id)
        result_by_step = {item.step.id: item for item in result.steps} if result else {}
        result_by_item = {move.step_item_id: move for step_result in result.steps for move in step_result.move_results} if result else {}
        start_row = row
        for step in action.steps:
            step_result = result_by_step.get(step.id)
            step_start = row
            items = step.items or [CompositeStepItem(kind="motion", distance_id=step.action1_distance_id)]
            for item in items:
                move = result_by_item.get(item.id)
                if item.kind == "delay":
                    axis_name = "延时"
                    action_label = f"{item.delay_ms:g} ms"
                else:
                    axis, _motion_action, distance = lookup.get(item.distance_id, (None, None, None))
                    axis_name = axis.name if axis else ""
                    action_label = _move_label(lookup, item.distance_id)
                ws4.append([
                    _composite_number(action, action_idx),
                    action.name,
                    step.name,
                    step.note,
                    axis_name,
                    action_label,
                    _fmt_time(move.total_time) if move else "",
                    _fmt_time(step_result.total_time) if step_result else "",
                    _fmt_time(result.total_time) if result else "",
                ])
                row += 1
            if row > step_start + 1:
                ws4.merge_cells(start_row=step_start, start_column=3, end_row=row - 1, end_column=3)
                ws4.merge_cells(start_row=step_start, start_column=4, end_row=row - 1, end_column=4)
                ws4.merge_cells(start_row=step_start, start_column=8, end_row=row - 1, end_column=8)
        if row > start_row + 1:
            ws4.merge_cells(start_row=start_row, start_column=9, end_row=row - 1, end_column=9)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        _auto_width(ws)

    for row in ws2.iter_rows(min_row=2, min_col=9, max_col=9):
        row[0].number_format = "0.000"
    for row in ws3.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "0.000"
    for row in ws4.iter_rows(min_row=2, min_col=7, max_col=9):
        for cell in row:
            cell.number_format = "0.000"

    wb.save(path)
