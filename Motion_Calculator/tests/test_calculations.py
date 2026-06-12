from motion_calculator.calculations import (
    calculate_forward,
    calculate_composite_motion,
    calculate_motion,
    calculate_t_profile,
    derive_axis,
    distance_to_register,
    register_to_motor_rpm,
    register_to_distance,
    register_to_tvmax_ms,
    register_to_velocity,
    recommend_for_target_time,
    tvmax_to_register,
    velocity_to_register,
)
from motion_calculator.models import Axis, AxisPoint, CompositeAction, CompositeStep, CompositeStepItem, DistanceCase, MotionCase, MotionParams, TProfileCase
from motion_calculator.models import AppState, MotionAction
from motion_calculator.io_excel import export_excel
from motion_calculator.project_io import load_project, save_project


def test_belt_axis_derived_values():
    axis = Axis(mechanism_type="linear_belt", motor_step_angle=1.8, microstep=256, pulley_teeth=20, pulley_pitch_mm=2)
    derived = derive_axis(axis)
    assert derived.count == 51200
    assert derived.unit_per_motor_rev == 40
    assert derived.microsteps_per_unit == 1280


def test_rotary_axis_derived_values():
    axis = Axis(mechanism_type="rotary", motor_step_angle=1.8, microstep=256, gear_ratio=2)
    derived = derive_axis(axis)
    assert derived.unit_name == "deg"
    assert derived.unit_per_motor_rev == 180


def test_forward_500mm_trapezoid():
    axis = Axis()
    case = MotionCase(distance_mm=500, vmax_mm_s=300, acc_mm_s2=1200, dec_mm_s2=1200)
    result = calculate_forward(axis, case)
    assert result.profile_type == "T型"
    assert round(result.total_time_s, 6) == round(300 / 1200 + 425 / 300 + 300 / 1200, 6)
    assert result.xtarget_register.value == 640000


def test_multistage_motion_has_segment_markers():
    axis = Axis()
    params = MotionParams(vmax=300, a1=800, a2=1000, amax=1200, dmax=1200, d2=1000, d1=800, v1=80, v2=180)
    result = calculate_motion(axis, 16_000_000, params, DistanceCase(distance=500))
    assert result.total_time > 0
    assert len(result.segments) >= 5
    assert result.vpeak > 0


def test_register_physical_roundtrip():
    axis = Axis()
    reg = velocity_to_register(axis, 16_000_000, 300)
    assert abs(register_to_velocity(axis, 16_000_000, reg.value) - 300) < 0.01
    xtarget = distance_to_register(axis, 500)
    assert register_to_distance(axis, xtarget.value) == 500


def test_prd_vmax_example_uses_motor_rpm_and_physical_speed():
    axis = Axis(mechanism_type="linear_belt", pulley_teeth=20, pulley_pitch_mm=2, microstep=256, motor_step_angle=1.8)
    rpm = register_to_motor_rpm(axis, 12_500_000, 50000)
    physical = register_to_velocity(axis, 12_500_000, 50000)
    assert abs(rpm - 43.8) < 0.2
    assert abs(physical - 29.2) < 0.2


def test_target_time_recommendation_balanced():
    axis = Axis()
    result = recommend_for_target_time(axis, distance_mm=500, target_time_s=2.0, strategy="balanced")
    assert abs(result.forward.total_time - 2.0) < 0.01
    assert result.acc_mm_s2 > 0
    assert result.dec_mm_s2 > 0
    assert result.params.tvmax_ms == 0


def test_target_time_recommendation_eight_low_impact_shape():
    axis = Axis()
    result = recommend_for_target_time(axis, distance_mm=500, target_time_s=2.0, strategy="eight_low_impact")
    params = result.params
    assert abs(result.forward.total_time - 2.0) < 0.05
    assert 0 < params.v1 < params.v2 < params.vmax
    assert params.a1 < params.a2
    assert params.amax < params.a2
    assert params.dmax < params.d2
    assert params.d1 < params.d2
    assert params.tvmax_ms == 0


def test_target_time_recommendation_eight_stall_guard_shape():
    axis = Axis()
    result = recommend_for_target_time(axis, distance_mm=500, target_time_s=2.0, strategy="eight_stall_guard")
    params = result.params
    assert params.a1 >= params.a2 >= params.amax
    assert params.d1 >= params.d2 >= params.dmax
    assert params.tvmax_ms == 0


def test_target_time_recommendation_respects_limits():
    axis = Axis()
    try:
        recommend_for_target_time(axis, distance_mm=500, target_time_s=0.5, strategy="eight_low_impact", vmax_limit=10, acc_limit=10, dec_limit=10)
    except ValueError as exc:
        assert "上限" in str(exc)
    else:
        raise AssertionError("expected recommendation to fail when limits are too low")


def test_t_profile_matches_prd_reference():
    axis = Axis()
    result = calculate_t_profile(axis, TProfileCase())
    assert abs(result.amax_register.raw - 27487.79) < 1
    assert abs(result.vmax_register.raw - 167772.16) < 1


def test_tvmax_register_conversion():
    assert abs(register_to_tvmax_ms(16_000_000, 1) - 0.032) < 1e-9
    assert abs(register_to_tvmax_ms(12_500_000, 1) - 0.04096) < 1e-9
    assert tvmax_to_register(16_000_000, 0.032).value == 1


def test_tvmax_zero_keeps_existing_motion_time():
    axis = Axis()
    base = MotionParams(vmax=300, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    with_tvmax_zero = MotionParams(
        vmax=base.vmax,
        tvmax_ms=0,
        a1=base.a1,
        a2=base.a2,
        amax=base.amax,
        dmax=base.dmax,
        d2=base.d2,
        d1=base.d1,
        v1=base.v1,
        v2=base.v2,
    )
    distance = DistanceCase(distance=500)
    assert calculate_motion(axis, 16_000_000, base, distance).total_time == calculate_motion(axis, 16_000_000, with_tvmax_zero, distance).total_time


def test_tvmax_lowers_vpeak_on_short_distance():
    axis = Axis()
    no_tvmax = MotionParams(vmax=300, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    with_tvmax = MotionParams(vmax=300, tvmax_ms=400, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    distance = DistanceCase(distance=40)
    base_result = calculate_motion(axis, 16_000_000, no_tvmax, distance)
    tvmax_result = calculate_motion(axis, 16_000_000, with_tvmax, distance)
    assert tvmax_result.vpeak < base_result.vpeak
    assert tvmax_result.vpeak < with_tvmax.vmax
    assert abs(tvmax_result.const_time - 0.4) < 1e-6


def test_tvmax_is_saved_and_loaded(tmp_path):
    axis = Axis()
    action = MotionAction(axis_id=axis.id, params=MotionParams(vmax=300, tvmax_ms=12.5))
    state = AppState(axes=[axis], actions=[action])
    path = tmp_path / "project.hmcalc"
    save_project(path, state)
    loaded = load_project(path)
    assert loaded.actions[0].params.tvmax_ms == 12.5


def test_composite_actions_are_saved_and_loaded(tmp_path):
    state = AppState(
        composite_actions=[
            CompositeAction(number="A01", name="上料", steps=[CompositeStep(name="STEP1", note="取料", delay_ms=100)]),
            CompositeAction(name="下料", steps=[CompositeStep(name="STEP1", delay_ms=200)]),
        ]
    )
    path = tmp_path / "project.hmcalc"
    save_project(path, state)
    loaded = load_project(path)
    assert [item.name for item in loaded.composite_actions] == ["上料", "下料"]
    assert loaded.composite_actions[0].number == "A01"
    assert loaded.composite_actions[0].steps[0].note == "取料"
    assert loaded.composite_actions[1].steps[0].delay_ms == 200


def test_legacy_composite_steps_migrate_to_one_action():
    state = AppState.from_dict({
        "composite_steps": [
            {"name": "STEP1", "delay_ms": 100},
            {"name": "STEP2", "delay_ms": 200},
        ]
    })
    assert len(state.composite_actions) == 1
    assert state.composite_actions[0].name == "组合动作1"
    assert len(state.composite_actions[0].steps) == 2


def test_composite_motion_sums_serial_and_maxes_parallel():
    x_axis = Axis(name="X")
    y_axis = Axis(name="Y")
    params = MotionParams(vmax=300, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    x1 = DistanceCase(distance=500, note="X1")
    y1 = DistanceCase(distance=100, note="Y1")
    x_action = MotionAction(axis_id=x_axis.id, params=params, distances=[x1])
    y_action = MotionAction(axis_id=y_axis.id, params=params, distances=[y1])
    x_time = calculate_motion(x_axis, 16_000_000, params, x1).total_time
    y_time = calculate_motion(y_axis, 16_000_000, params, y1).total_time

    result = calculate_composite_motion(
        [x_axis, y_axis],
        [x_action, y_action],
        [
            CompositeStep(name="STEP1", action1_distance_id=x1.id, action2_distance_id=y1.id),
            CompositeStep(name="STEP2", delay_ms=250),
        ],
    )

    assert result.steps[0].total_time == max(x_time, y_time)
    assert abs(result.total_time - (max(x_time, y_time) + 0.25)) < 1e-9


def test_legacy_distances_migrate_to_axis_points():
    axis = Axis(name="X")
    distance = DistanceCase(distance=125, note="取料位")
    state = AppState(axes=[axis], actions=[MotionAction(axis_id=axis.id, distances=[distance])])
    state.migrate_coordinate_points()

    migrated = state.actions[0].distances[0]
    assert len(state.axes[0].points) == 2
    assert state.axes[0].points[0].position_mm == 0
    assert migrated.start_point_id == state.axes[0].points[0].id
    assert migrated.end_point_id == state.axes[0].points[1].id
    assert migrated.distance == 125


def test_signed_coordinate_distance_uses_absolute_motion_time():
    zero = AxisPoint(name="零位", position_mm=0)
    negative = AxisPoint(name="负位", position_mm=-100)
    axis = Axis(name="X", points=[zero, negative])
    params = MotionParams(vmax=300, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    distance = DistanceCase(distance=-100, note="反向", start_point_id=zero.id, end_point_id=negative.id)
    action = MotionAction(axis_id=axis.id, params=params, distances=[distance])
    step = CompositeStep(name="STEP1", items=[CompositeStepItem(kind="motion", axis_id=axis.id, distance_id=distance.id)])

    result = calculate_composite_motion([axis], [action], [step])

    assert distance_to_register(axis, distance.distance).value < 0
    assert result.total_time == calculate_motion(axis, 16_000_000, params, DistanceCase(distance=100)).total_time
    assert result.steps[0].move_results[0].distance == -100


def test_composite_step_items_support_delay_after_parallel_motion():
    axis = Axis(name="X")
    params = MotionParams(vmax=300, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    distance = DistanceCase(distance=100, note="P1")
    action = MotionAction(axis_id=axis.id, params=params, distances=[distance])
    motion_time = calculate_motion(axis, 16_000_000, params, distance).total_time
    step = CompositeStep(
        name="STEP1",
        items=[
            CompositeStepItem(kind="motion", axis_id=axis.id, distance_id=distance.id),
            CompositeStepItem(kind="delay", delay_ms=200),
        ],
    )

    result = calculate_composite_motion([axis], [action], [step])

    assert abs(result.steps[0].total_time - (motion_time + 0.2)) < 1e-9
    assert result.steps[0].delay_time == 0.2


def test_timer_axis_secondary_action_participates_in_parallel_step_time():
    timer = Axis(name="Timer", mechanism_type="timer")
    timer_case = DistanceCase(number="T1", duration_s=0.75, note="等待")
    timer_action = MotionAction(axis_id=timer.id, distances=[timer_case])
    step = CompositeStep(
        name="STEP1",
        items=[CompositeStepItem(kind="motion", axis_id=timer.id, distance_id=timer_case.id)],
    )

    result = calculate_composite_motion([timer], [timer_action], [step])

    assert result.total_time == 0.75
    assert result.steps[0].move_results[0].kind == "timer"


def test_excel_export_supports_current_app_state(tmp_path):
    axis = Axis(name="X")
    distance = DistanceCase(distance=120, note="取料位")
    params = MotionParams(vmax=300, a1=1200, a2=1200, amax=1200, dmax=1200, d2=1200, d1=1200, v1=100, v2=200)
    action = MotionAction(name="X 参数", axis_id=axis.id, params=params, distances=[distance])
    composite = CompositeAction(name="上料", steps=[CompositeStep(name="STEP1", note="并行取料", action1_distance_id=distance.id)])
    state = AppState(project_name="测试项目", axes=[axis], actions=[action], composite_actions=[composite])

    path = tmp_path / "export.xlsx"
    export_excel(path, state)

    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    assert "项目与单轴参数" in wb.sheetnames
    assert "距离坐标" in wb.sheetnames
    assert "一级动作时间" in wb.sheetnames
    assert "组合STEP列表" in wb.sheetnames
    assert wb["项目与单轴参数"]["B2"].value == "测试项目"
    assert wb["组合STEP列表"]["D2"].value == "并行取料"
